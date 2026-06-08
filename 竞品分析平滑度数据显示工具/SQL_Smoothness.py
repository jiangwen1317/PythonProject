import pandas as pd
import os
import PySimpleGUI as sg
import matplotlib.pyplot as plt
import random
import glob
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
import zipfile
import re
import shutil
import sys

# 配置项与Excel列的映射 (使用小写+下划线归一化匹配)
CONFIG_MAP = {
    'insert_burn': 1,
    'insert_load': 2,
    'updata_burn': 3,
    'updata_load': 4,
    'delect_burn': 5,
    'delect_load': 6
}


def extract_nested_zips(start_dir, max_iterations=2):
    """深度解压指定目录下的所有 ZIP 文件，并删除原始压缩包（优化：使用列表推导式）"""
    print(f"【开始】深度解压 {start_dir} 下的所有 ZIP 文件，并删除压缩包...")
    for iteration in range(1, max_iterations + 1):
        zip_files = [
            os.path.join(root, file)
            for root, _, files in os.walk(start_dir)
            for file in files
            if file.lower().endswith('.zip')
        ]

        if not zip_files:
            print(f"【完成】第 {iteration} 轮：未发现 ZIP 文件，停止解压。")
            break

        print(f"【解压】第 {iteration} 轮：发现 {len(zip_files)} 个 ZIP 文件...")
        extracted = False
        for zip_path in zip_files:
            try:
                extract_dir = os.path.dirname(zip_path)
                base_name = os.path.basename(zip_path)
                with zipfile.ZipFile(zip_path, 'r') as z:
                    z.extractall(extract_dir)
                print(f"          已解压: {base_name}")
                os.remove(zip_path)
                print(f"          已删除: {base_name}")
                extracted = True
            except Exception as e:
                print(f"          解压失败: {os.path.basename(zip_path)} - {e}")

        if not extracted:
            print(f"【完成】第 {iteration} 轮后无新文件解压，结束。")
            break
    print("【清理】所有可解压的 ZIP 文件已处理并删除。")


# 随机分配颜色代码
def getcolor(n):
    if n <= 0: return []
    ans = []
    for i in range(n):
        r = random.randint(50, 200)
        g = random.randint(50, 200)
        b = random.randint(50, 200)
        ans.append(f'#{r:02x}{g:02x}{b:02x}')
    return ans


def apply_scaling(data, val):
    """应用谐波平均缩放逻辑，严格匹配原版逻辑"""
    val = max(1, int(val))
    tempx = []
    temsetpy = []

    # 按照原版循环逻辑，即使val=1也通过公式计算
    for x in range(len(data) // val):
        temp_sum = 0
        tempx.append(x)
        for j in range(val):
            val_in_data = float(data[(x * val) + j])
            if val_in_data == 0:
                val_in_data = 1.0  # 原版逻辑：0变为1
            temp_sum += 1000.0 / val_in_data

        if temp_sum == 0:
            tempvaly = 0
        else:
            tempvaly = val * 1000.0 / temp_sum
        temsetpy.append(tempvaly)
    return tempx, temsetpy


def process_folder(root_path, scale_val, filter_keywords=None):
    """主处理逻辑"""
    output_images = []

    # 关键词预处理
    keywords = []
    if filter_keywords:
        keywords = [k.strip().lower() for k in filter_keywords.replace('，', ',').split(',') if k.strip()]

    # 遍历一级目录 (读写方式)
    for lvl1 in os.listdir(root_path):
        lvl1_path = os.path.join(root_path, lvl1)
        if not os.path.isdir(lvl1_path): continue

        # 遍历二级目录 (配置名)
        for lvl2 in os.listdir(lvl1_path):
            lvl2_path = os.path.join(lvl1_path, lvl2)
            if not os.path.isdir(lvl2_path): continue

            # 归一化目录名：转小写，并将中划线替换为下划线，以增加匹配兼容性
            lvl2_norm = lvl2.lower().replace('-', '_')

            if lvl2_norm not in CONFIG_MAP:
                print(f"跳过未知配置目录: {lvl2} (归一化后仍未匹配: {lvl2_norm})")
                continue

            xlsx_files = glob.glob(os.path.join(lvl2_path, "*.xlsx"))
            if not xlsx_files: continue

            # 文件筛选逻辑 (模糊匹配)
            if keywords:
                filtered_files = []
                for f in xlsx_files:
                    f_name_lower = os.path.basename(f).lower()
                    if any(k in f_name_lower for k in keywords):
                        filtered_files.append(f)
                xlsx_files = filtered_files

            if not xlsx_files:
                print(f"目录 {lvl2} 下未发现符合筛选条件的文件")
                continue

            print(f"正在处理: {lvl1} -> {lvl2} (包含 {len(xlsx_files)} 个文件)")

            plt.close('all')
            fig = plt.figure(figsize=(16, 9), dpi=100)
            ax = fig.add_subplot(1, 1, 1)

            col_idx = CONFIG_MAP[lvl2_norm]
            colors = getcolor(len(xlsx_files))

            max_y = 0
            max_x = 0

            for idx, f_path in enumerate(xlsx_files):
                f_name = os.path.splitext(os.path.basename(f_path))[0]  # 去掉扩展名
                try:
                    # 读取第3行到20002行 (skiprows=2, nrows=20000)
                    df = pd.read_excel(f_path, header=None, skiprows=2, nrows=20000, usecols=[col_idx])
                    data_raw = df.iloc[:, 0].tolist()

                    x_data, y_data = apply_scaling(data_raw, scale_val)

                    ax.plot(x_data, y_data, color=colors[idx], label=f_name)

                    if x_data: max_x = max(max_x, max(x_data))
                    if y_data: max_y = max(max_y, max(y_data))
                except Exception as e:
                    print(f"读取文件 {f_name} 出错: {e}")

            ax.set_title(f"{lvl1}-{lvl2}-平滑度曲线")
            ax.set_xlabel("采样点")
            ax.set_ylabel("速度")
            ax.set_xlim(left=0)
            ax.set_ylim(bottom=0)
            if max_y > 0:
                ax.set_ylim(top=max_y * 1.1)

            ax.legend(bbox_to_anchor=(1.01, 1), loc='upper left', fontsize='small')
            plt.tight_layout()
            plt.grid(alpha=0.3)

            # 保存临时图片
            img_name = f"{lvl1}_{lvl2}_smoothness.png"
            img_path = os.path.join(root_path, img_name)
            fig.savefig(img_path)
            output_images.append((lvl1, lvl2, img_path))

    return output_images


def create_excel_report(root_path, image_data):
    """创建Excel报告并插入图片"""
    report_path = os.path.join(root_path, "平滑度分析汇总报告.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总图表"

    ws.append(["读写方式", "配置名", "平滑度曲线图"])

    current_row = 2
    for lvl1, lvl2, img_path in image_data:
        ws.cell(row=current_row, column=1, value=lvl1)
        ws.cell(row=current_row, column=2, value=lvl2)

        img = OpenpyxlImage(img_path)
        # 适当调整图片大小以适应单元格 (可选)
        img.width = 600
        img.height = 350

        # 插入图片到C列
        anchor = f'C{current_row}'
        ws.add_image(img, anchor)

        # 设置行高，以便容纳图片
        ws.row_dimensions[current_row].height = 270
        current_row += 1

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 90

    wb.save(report_path)
    print(f"报告已生成: {report_path}")


def process_logs_to_excel(root_path):
    """从日志中提取数据并填入 Excel 模板"""
    # 1. 深度解压
    extract_nested_zips(root_path)

    # 2. 配置映射定义
    config_targets = {
        "1-CMDQ_In_burn.json": "WSpeed_In:",
        "2-CMDQ_In_burn&load.json": "WSpeed_In_Load:",
        "1-PACK_In_burn.json": "WSpeed_In:",
        "2-PACK_In_burn&load.json": "WSpeed_In_Load:",
        "1-CMD25_CMD12_In_burn.json": "WSpeed_In:",
        "2-CMD25_CMD12_In_burn&load.json": "WSpeed_In_Load:",
        "1-CMD23_CMD25_In_burn.json": "WSpeed_In:",
        "2-CMD23_CMD25_In_burn&load.json": "WSpeed_In_Load:"
    }

    # 精简版路径识别：兼容源码运行与打包后的 EXE 环境
    base_dir = os.path.dirname(os.path.realpath(sys.executable if getattr(sys, 'frozen', False) else __file__))
    template_path = os.path.join(base_dir, "SQL空模板.xlsx")

    if not os.path.exists(template_path):
        print(f"【错误】未找到模板文件：{template_path}")
        return 0

    excel_count = 0
    print(f"【开始】正在扫描目录 {root_path} 下的 TXT 日志文件...")

    for root, _, files in os.walk(root_path):
        for file in files:
            if file.lower().endswith('.txt'):
                txt_path = os.path.join(root, file)
                try:
                    # 读取日志内容
                    with open(txt_path, 'r', encoding='utf-8', errors='ignore') as f:
                        content = f.read()

                    # 提取 Config_Name
                    config_match = re.search(r'Config_Name\s*:\s*([\w\.\-\&]+)', content)
                    if not config_match:
                        continue

                    cfg_name = config_match.group(1).strip()
                    if cfg_name in config_targets:
                        keyword = config_targets[cfg_name]
                        # 查找所有匹配关键字后的数字
                        pattern = rf"{re.escape(keyword)}\s*([\d\.]+)"
                        values = re.findall(pattern, content)

                        if not values:
                            print(f"      [警告] 文件 {file} 匹配到配置 {cfg_name} 但未找到关键字 {keyword}")
                            continue

                        # 生成输出文件路径
                        output_name = f"{cfg_name}.xlsx"
                        output_path = os.path.join(root, output_name)

                        # 复制模板并写入
                        shutil.copy(template_path, output_path)
                        wb = openpyxl.load_workbook(output_path)

                        # 获取 sheet1，如果不存在则使用活跃 sheet
                        ws = wb['sheet1'] if 'sheet1' in wb.sheetnames else wb.active

                        # 获取板子号：当前文件夹(root)的上2级文件夹名字
                        try:
                            parent_dir = os.path.dirname(root)
                            grandparent_dir = os.path.dirname(parent_dir)
                            board_id = os.path.basename(grandparent_dir)
                            ws.cell(row=3, column=1, value=board_id)
                        except Exception as e:
                            print(f"      [警告] 提取板子号失败: {e}")

                        # B列(2) 对应 WSpeed_In:，C列(3) 对应 WSpeed_In_Load:
                        col_idx = 2 if keyword == "WSpeed_In:" else 3

                        for idx, val in enumerate(values):
                            try:
                                ws.cell(row=3 + idx, column=col_idx, value=float(val))
                            except ValueError:
                                ws.cell(row=3 + idx, column=col_idx, value=val)

                        wb.save(output_path)
                        print(f"      [成功] 已生成: {output_path} (提取到 {len(values)} 条数据)")
                        excel_count += 1

                except Exception as e:
                    print(f"      [失败] 处理文件 {file} 时出错: {e}")

    return excel_count


if __name__ == '__main__':
    plt.rcParams['font.sans-serif'] = ['SimHei']
    plt.rcParams['axes.unicode_minus'] = False

    # 定义选项卡1：日志解析
    tab1_layout = [
        [sg.Text('说明：从原始日志中提取 WSpeed 数据并生成 Excel 模板。', text_color='yellow')],
        [sg.Button('一键提取日志数据', size=(20, 1), key='-GEN_EXCEL-', button_color=('white', '#28a745'))]
    ]

    # 定义选项卡2：图表生成
    tab2_layout = [
        [sg.Text('说明：遍历一级/二级目录下的 Excel 文件生成平滑度曲线。')],
        [sg.Text('文件筛选（可选，支持逗号分隔/模糊匹配）：')],
        [sg.Input(key='-FILTER-', size=(60, 1))],
        [sg.Text('设置缩放倍数 (1表示不缩放)：')],
        [sg.Slider(range=(1, 100), key='-SCALE-', default_value=1, orientation='h', size=(30, 15))],
        [sg.Button('生成平滑度分析报告', size=(20, 1), key='-GEN_CHART-', button_color=('white', '#007bff'))]
    ]

    layout = [
        [sg.Text('请选择数据根目录：')],
        [sg.Input(key='-ROOT-', size=(50, 1)), sg.FolderBrowse('选择文件夹')],
        [sg.TabGroup([[
            sg.Tab('日志转Excel', tab1_layout),
            sg.Tab('Excel转图表', tab2_layout)
        ]], border_width=1)],
        [sg.HorizontalSeparator()],
        [sg.Button('退出', size=(10, 1))]
    ]

    window = sg.Window("SQL_平滑度批量分析工具_V2.1.0 ", layout)

    while True:
        event, values = window.read()

        if event in (sg.WIN_CLOSED, '退出'):
            break

        root_dir = values['-ROOT-']
        if not root_dir or not os.path.exists(root_dir):
            if event != sg.WIN_CLOSED:
                sg.popup_error("请先选择有效的根目录！")
            continue

        if event == '-GEN_CHART-':
            scale_factor = int(values['-SCALE-'])
            filter_str = values['-FILTER-']

            sg.popup_quick_message("正在生成图表，请稍候...", background_color='gray', text_color='white')

            try:
                images = process_folder(root_dir, scale_factor, filter_str)
                if not images:
                    sg.popup("未发现符合结构的测试数据（一级/二级目录格式）。")
                else:
                    create_excel_report(root_dir, images)
                    sg.popup(f"图表处理完成！\n共生成 {len(images)} 张图片。\n报告已保存至根目录。")
            except Exception as e:
                sg.popup_error(f"处理图表时发生错误: {e}")

        elif event == '-GEN_EXCEL-':
            sg.popup_quick_message("正在从日志提取数据，请稍候...", background_color='gray', text_color='white')
            try:
                count = process_logs_to_excel(root_dir)
                if count > 0:
                    sg.popup(f"日志提取完成！\n共生成 {count} 个 Excel 文件。")
                else:
                    sg.popup("未发现符合条件的日志文件或配置名。")
            except Exception as e:
                sg.popup_error(f"处理日志时发生错误: {e}")

    window.close()
