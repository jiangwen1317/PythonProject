import pandas as pd
import os
import sys
import PySimpleGUI as sg
import matplotlib.pyplot as plt
import random
import glob
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
import zipfile
import re
import shutil


# 统一配置映射：定义指标名称、日志关键字和Excel列索引 (C-F列对应 3-6)
MASTER_CONFIG = {
    '0% WriteSpeed':  {'log_key': 'dwCurrentWriteMBps_0:',   'col': 3},
    '50% WriteSpeed': {'log_key': 'dwCurrentWriteMBps_50:',  'col': 4},
    '50% ReadSpeed':  {'log_key': 'dwCurrentReadMBps_50:',   'col': 5},
    '100% ReadSpeed': {'log_key': 'dwCurrentReadMBps_100:',  'col': 6}
}

# 绘图使用的配置映射
CONFIG_MAP = {name: info['col'] for name, info in MASTER_CONFIG.items()}


def extract_nested_zips(start_dir, max_iterations=2):
    """深度解压指定目录下的所有 ZIP 文件，并删除原始压缩包（优化：使用列表推导式）"""
    print(f"【开始】深度解压 {start_dir} 下的所有 ZIP 文件，并删除压缩包...")
    for iteration in range(1, max_iterations + 1):
        zip_files = [
            os.path.join(root, file)
            for root, _, files in os.walk(start_dir)
            for file in files
            if file.lower().endswith('.zip')
        ]

        if not zip_files:
            print(f"【完成】第 {iteration} 轮：未发现 ZIP 文件，停止解压。")
            break

        print(f"【解压】第 {iteration} 轮：发现 {len(zip_files)} 个 ZIP 文件...")
        extracted = False
        for zip_path in zip_files:
            try:
                extract_dir = os.path.dirname(zip_path)
                base_name = os.path.basename(zip_path)
                with zipfile.ZipFile(zip_path, 'r') as z:
                    z.extractall(extract_dir)
                print(f"          已解压: {base_name}")
                os.remove(zip_path)
                print(f"          已删除: {base_name}")
                extracted = True
            except Exception as e:
                print(f"          解压失败: {os.path.basename(zip_path)} - {e}")

        if not extracted:
            print(f"【完成】第 {iteration} 轮后无新文件解压，结束。")
            break
    print("【清理】所有可解压的 ZIP 文件已处理并删除。")


def process_logs_to_excel(root_path):
    """从日志中提取RW平滑度数据并填入 Excel 模板"""
    # 1. 深度解压
    extract_nested_zips(root_path)

    # 2. 获取指标映射 (从 MASTER_CONFIG 统一获取)
    field_map = {info['log_key']: info['col'] for info in MASTER_CONFIG.values()}

    # 3. 获取模板路径（同级目录查找）
    base_dir = os.path.dirname(os.path.realpath(sys.executable if getattr(sys, 'frozen', False) else __file__))
    template_path = os.path.join(base_dir, "IO_BEKK空模板.xlsx")

    if not os.path.exists(template_path):
        print(f"【错误】未找到模板文件：{template_path}")
        return 0

    excel_count = 0
    print(f"【开始】正在扫描目录 {root_path} 下的 TXT 日志文件...")

    # 4. 遍历所有txt文件
    for root, _, files in os.walk(root_path):
        for file in files:
            if file.lower().endswith('.txt'):
                txt_path = os.path.join(root, file)
                try:
                    # 读取日志内容
                    with open(txt_path, 'r', encoding='utf-8', errors='ignore') as f:
                        content = f.read()

                    # 提取配置名
                    config_match = re.search(r'Config_Name:(.+)', content)
                    if not config_match:
                        print(f"      [跳过] 文件 {file} 未找到 Config_Name 字段")
                        continue

                    config_name = config_match.group(1).strip()
                    # 去掉.json后缀
                    if config_name.endswith('.json'):
                        config_name = config_name[:-5]

                    # 提取板子号：当前文件夹(root)的上2级文件夹名字
                    try:
                        parent_dir = os.path.dirname(root)
                        grandparent_dir = os.path.dirname(parent_dir)
                        board_id = os.path.basename(grandparent_dir)
                    except Exception:
                        board_id = "Unknown"

                    # 创建数据存储字典
                    data_dict = {field: [] for field in field_map.keys()}
                    has_data = False

                    # 提取每个字段的数据
                    for field_name, col_idx in field_map.items():
                        # 使用findall提取所有出现的数据
                        pattern = rf"{re.escape(field_name)}\s*([\d,\s]+?)(?=\n|\r|$)"
                        matches = re.findall(pattern, content, re.MULTILINE)

                        all_numbers = []
                        for match_str in matches:
                            # 从每个匹配中提取所有数字
                            numbers = re.findall(r'\d+', match_str)
                            all_numbers.extend(numbers)

                        data_dict[field_name] = all_numbers
                        if all_numbers:
                            has_data = True
                            print(f"        提取 {field_name} -> {len(all_numbers)} 个数据")

                    # 如果没有提取到任何数据，跳过该文件
                    if not has_data:
                        continue

                    # 生成输出文件名（使用配置名）
                    output_name = f"{config_name}.xlsx"
                    output_path = os.path.join(root, output_name)

                    # 复制模板并写入
                    shutil.copy(template_path, output_path)
                    wb = openpyxl.load_workbook(output_path)
                    ws = wb.active

                    # 写入板子号到A列第2行
                    ws.cell(row=2, column=1, value=board_id)

                    # 写入配置名到B列第2行
                    ws.cell(row=2, column=2, value=config_name)

                    # 写入数据到对应列（从第2行开始）
                    for field_name, col_idx in field_map.items():
                        values = data_dict[field_name]
                        for idx, val in enumerate(values):
                            try:
                                ws.cell(row=2 + idx, column=col_idx, value=float(val))
                            except ValueError:
                                ws.cell(row=2 + idx, column=col_idx, value=val)

                    wb.save(output_path)
                    total_data = sum(len(v) for v in data_dict.values())
                    print(f"      [成功] 已生成: {output_path} (提取到 {total_data} 条数据)")
                    excel_count += 1

                except Exception as e:
                    print(f"      [失败] 处理文件 {file} 时出错: {e}")

    return excel_count


# 配置项与Excel列的映射已统一移至 MASTER_CONFIG
# 随机分配颜色代码
def get_colors(n):
    if n <= 0: return []
    ans = []
    for i in range(n):
        r = random.randint(30, 200)
        g = random.randint(30, 200)
        b = random.randint(30, 200)
        ans.append(f'#{r:02x}{g:02x}{b:02x}')
    return ans


def apply_scaling(data, val):
    """应用谐波平均缩放逻辑，严格匹配原版逻辑"""
    val = max(1, int(val))
    tempx = []
    temsetpy = []

    # 按照原版循环逻辑，即使val=1也通过公式计算
    for x in range(len(data) // val):
        temp_sum = 0
        tempx.append(x)
        for j in range(val):
            try:
                val_in_data = float(data[(x * val) + j])
            except (ValueError, TypeError, IndexError):
                val_in_data = 1.0

            if val_in_data == 0:
                val_in_data = 1.0  # 原版逻辑：0变为1
            temp_sum += 1000.0 / val_in_data

        if temp_sum == 0:
            tempvaly = 0
        else:
            tempvaly = val * 1000.0 / temp_sum
        temsetpy.append(tempvaly)
    return tempx, temsetpy


def process_folder_to_charts(root_path, scale_val, line_type, filter_keywords=None):
    """
    扫描目录下所有Excel，为每个配置生成一张汇总对比图
    """
    folder_name = os.path.basename(root_path)
    print(f"\n【正在扫描文件夹】: {folder_name}")

    xlsx_files = glob.glob(os.path.join(root_path, "*.xlsx"))
    # 排除生成的报告文件
    xlsx_files = [f for f in xlsx_files if "平滑度分析" not in f]

    # 文件筛选逻辑 (模糊匹配)
    if filter_keywords:
        keywords = [k.strip().lower() for k in filter_keywords.replace('，', ',').split(',') if k.strip()]
        if keywords:
            filtered_files = []
            for f in xlsx_files:
                f_name_lower = os.path.basename(f).lower()
                if any(k in f_name_lower for k in keywords):
                    filtered_files.append(f)
            xlsx_files = filtered_files

    if not xlsx_files:
        print(f"  [提示] 文件夹 {folder_name} 下未发现符合条件的Excel文件，跳过。")
        return []

    print(f"  [发现文件] 共 {len(xlsx_files)} 个符合条件的Excel文件。")
    output_images = []

    # 线条样式
    plot_style = '-' if line_type == '线条' else '.'

    # 为每个配置生成一张图
    for config_name, col_idx in CONFIG_MAP.items():
        print(f"    -> 正在生成配置图表: {config_name}...", end="", flush=True)
        plt.close('all')
        fig = plt.figure(figsize=(12, 7), dpi=100)
        ax = fig.add_subplot(1, 1, 1)

        colors = get_colors(len(xlsx_files))
        max_y = 0

        has_data = False
        for idx, f_path in enumerate(xlsx_files):
            f_name = os.path.splitext(os.path.basename(f_path))[0]
            try:
                # 读取数据 (pandas的usecols从0开始，openpyxl的列索引从1开始，所以需要-1)
                df = pd.read_excel(f_path, header=None, skiprows=1, usecols=[col_idx - 1])
                data_raw = df.iloc[:, 0].dropna().tolist()

                if not data_raw:
                    continue

                x_data, y_data = apply_scaling(data_raw, scale_val)

                if y_data:
                    ax.plot(x_data, y_data, plot_style, color=colors[idx], label=f_name)
                    max_y = max(max_y, max(y_data))
                    has_data = True
            except Exception as e:
                print(f"\n      [错误] 处理文件 {f_name} 时出错: {e}")

        if not has_data:
            print(" 无有效数据。")
            continue

        # 设置图表格式
        ax.set_title(f"{folder_name}-{config_name}-平滑度曲线", fontsize=12)
        ax.set_xlabel("采样点")
        ax.set_ylabel("速度")
        ax.set_xlim(left=0)
        ax.set_ylim(bottom=0)
        if max_y > 0:
            ax.set_ylim(top=max_y * 1.1)

        ax.legend(bbox_to_anchor=(1.01, 1), loc='upper left', fontsize='xx-small')
        plt.grid(alpha=0.3)
        plt.tight_layout()

        # 保存图片
        img_name = f"{folder_name}_{config_name.replace('%', 'pct').replace(' ', '_')}.png"
        img_path = os.path.join(root_path, img_name)
        fig.savefig(img_path)
        output_images.append((config_name, img_path))
        print(" 完成")

    return output_images


def create_excel_report(root_path, all_folder_results):
    """创建汇总报告，文件夹横向排布"""
    print(f"\n【正在汇总报告】...")
    report_path = os.path.join(root_path, "平滑度分析批量汇总报告.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "矩阵对比图表"

    config_names = ['0% WriteSpeed', '50% WriteSpeed', '50% ReadSpeed', '100% ReadSpeed']

    # 第一列：配置名称
    for idx, name in enumerate(config_names):
        ws.cell(row=idx + 2, column=1, value=name)
    ws.column_dimensions['A'].width = 15

    # 不同文件夹横向排列
    for f_idx, (folder_name, image_data) in enumerate(all_folder_results):
        col_idx = f_idx + 2
        col_letter = get_column_letter(col_idx)

        # 第一行：文件夹名称
        ws.cell(row=1, column=col_idx, value=folder_name)
        ws.column_dimensions[col_letter].width = 90

        # 将图片数据转为字典方便查找
        img_dict = {name: path for name, path in image_data}

        for r_idx, config_name in enumerate(config_names):
            current_row = r_idx + 2
            ws.row_dimensions[current_row].height = 300

            if config_name in img_dict:
                img_path = img_dict[config_name]
                try:
                    img = OpenpyxlImage(img_path)
                    img.width = 640
                    img.height = 360
                    ws.add_image(img, f"{col_letter}{current_row}")
                except Exception as e:
                    print(f"    [警告] 插入图片 {img_path} 失败: {e}")

    wb.save(report_path)
    print(f"【报告生成成功】: {report_path}")

    # 清理临时图片
    print("【正在清理临时文件】...")
    for _, image_data in all_folder_results:
        for _, img_path in image_data:
            try:
                if os.path.exists(img_path): os.remove(img_path)
            except:
                pass

    return report_path


def main():
    plt.rcParams['font.sans-serif'] = ['SimHei']
    plt.rcParams['axes.unicode_minus'] = False

    # 定义选项卡1：日志解析
    tab1_layout = [
        [sg.Text('说明：从原始日志中提取 dwCurrentWriteMBps/dwCurrentReadMBps 数据并生成 Excel 模板。',
                 text_color='yellow')],
        [sg.Button('一键提取日志数据', size=(20, 1), key='-GEN_EXCEL-', button_color=('white', '#28a745'))]
    ]

    # 定义选项卡2：图表生成
    tab2_layout = [
        [sg.Text('说明：扫描目录及子目录下的 Excel 文件生成平滑度曲线对比图。')],
        [sg.Text('文件筛选（可选，支持逗号分隔/模糊匹配）：')],
        [sg.Input(key='-FILTER-', size=(60, 1))],
        [sg.Text('设置缩放倍数 (1表示不缩放)：')],
        [sg.Slider(range=(1, 100), key='-SCALE-', default_value=1, orientation='h', size=(30, 15))],
        [sg.Text('选择线条类型：')],
        [sg.Combo(['线条', '散点'], key='-STYLE-', default_value='线条', size=(10, 1))],
        [sg.Button('生成平滑度分析报告', size=(20, 1), key='-GEN_CHART-', button_color=('white', '#007bff'))]
    ]

    layout = [
        [sg.Text('请选择数据根目录：')],
        [sg.Input(key='-ROOT-', size=(50, 1)), sg.FolderBrowse('选择文件夹')],
        [sg.TabGroup([[
            sg.Tab('日志转Excel', tab1_layout),
            sg.Tab('Excel转图表', tab2_layout)
        ]], border_width=1)],
        [sg.HorizontalSeparator()],
        [sg.Button('退出', size=(10, 1))]
    ]

    window = sg.Window("RW_平滑度批量分析工具_V2.1.0", layout)

    while True:
        event, values = window.read()

        if event in (sg.WIN_CLOSED, '退出'):
            break

        root_dir = values['-ROOT-']
        if not root_dir or not os.path.exists(root_dir):
            if event not in (sg.WIN_CLOSED, '退出'):
                sg.popup_error("请先选择有效的根目录！")
            continue

        if event == '-GEN_CHART-':
            scale_val = int(values['-SCALE-'])
            line_style = values['-STYLE-']
            filter_str = values['-FILTER-']

            # 识别目标文件夹：当前目录，以及包含Excel的子目录
            target_folders = []
            if glob.glob(os.path.join(root_dir, "*.xlsx")):
                target_folders.append(root_dir)

            for item in os.listdir(root_dir):
                sub_p = os.path.join(root_dir, item)
                if os.path.isdir(sub_p):
                    if glob.glob(os.path.join(sub_p, "*.xlsx")):
                        target_folders.append(sub_p)

            if not target_folders:
                sg.popup_error("选择的目录及其子目录下未发现 .xlsx 文件。")
                continue

            sg.popup_quick_message(f"正在处理 {len(target_folders)} 个文件夹，请稍候...", background_color='gray',
                                   text_color='white')

            try:
                all_results = []
                for folder in target_folders:
                    images = process_folder_to_charts(folder, scale_val, line_style, filter_str)
                    if images:
                        all_results.append((os.path.basename(folder), images))

                if not all_results:
                    sg.popup("处理完毕，但未产生有效图表。")
                else:
                    report_path = create_excel_report(root_dir, all_results)
                    sg.popup(f"批量处理完成！\n共分析了 {len(all_results)} 个文件夹。\n报告已生成：\n{report_path}")
            except Exception as e:
                sg.popup_error(f"发生错误: {e}")
                import traceback
                traceback.print_exc()

        elif event == '-GEN_EXCEL-':
            sg.popup_quick_message("正在从日志提取数据，请稍候...", background_color='gray', text_color='white')
            try:
                count = process_logs_to_excel(root_dir)
                if count > 0:
                    sg.popup(f"日志提取完成！\n共生成 {count} 个 Excel 文件。")
                else:
                    sg.popup("未发现符合条件的日志文件。")
            except Exception as e:
                sg.popup_error(f"处理日志时发生错误: {e}")
                import traceback
                traceback.print_exc()

    window.close()


if __name__ == '__main__':
    main()
