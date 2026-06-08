import os
import re
import time

import tkinter as tk
from tkinter import filedialog, messagebox
import pandas as pd
from openpyxl import Workbook, load_workbook

# 用户选择的数据文件夹路径
FILE = ""


def file_name(file_dir, source_type):
    """递归查找指定扩展名的所有文件
    
    Args:
        file_dir: 目录路径
        source_type: 文件扩展名（如 '.txt', '.zip'）
        
    Returns:
        list: 文件路径列表；若未找到则返回 0
    """
    file_list = []
    for root, dirs, files in os.walk(file_dir):
        for file in files:
            if os.path.splitext(file)[1] == source_type:
                file_list.append(os.path.join(root, file))

    return file_list if file_list else 0


def save_board_error_info(worksheet, row, col, content):
    """写入数据到指定单元格，并返回下一列位置"""
    worksheet.cell(row, col).value = content
    return col + 1


def Pop_Up_Window_1():
    """显示文件夹选择对话框，结果存入全局变量 FILE"""
    global FILE

    window = tk.Tk()
    window.title('EVT性能数据汇总工具_V2.0.8')
    window.geometry('550x200')

    folder_path = tk.StringVar()

    # 文件夹选择
    frame1 = tk.Frame(window)
    frame1.pack(pady=10, padx=10, fill=tk.X)
    tk.Button(frame1, text='请选择你的文件路径', command=lambda: folder_path.set(
        filedialog.askdirectory() or folder_path.get())).pack(side=tk.LEFT)

    # 显示选择的路径
    frame2 = tk.Frame(window)
    frame2.pack(pady=5, padx=10, fill=tk.X)
    tk.Label(frame2, text='你选择数据库所在文件夹是：').pack(side=tk.LEFT)
    tk.Label(frame2, textvariable=folder_path, bg='white', fg='black',
             anchor=tk.W, relief=tk.SUNKEN, width=50).pack(side=tk.LEFT, fill=tk.X, expand=True)

    # 按钮
    frame3 = tk.Frame(window)
    frame3.pack(pady=10)

    def on_confirm():
        global FILE
        FILE = folder_path.get()
        print(FILE)
        window.destroy()

    def on_exit():
        window.destroy()

    tk.Button(frame3, text='确认', command=on_confirm).pack(side=tk.LEFT, padx=10)
    tk.Button(frame3, text='退出', command=on_exit).pack(side=tk.LEFT, padx=10)

    window.protocol("WM_DELETE_WINDOW", on_exit)
    window.mainloop()


def process_fields(file_path, keys):
    """汇总关键字段写入Excel"""
    temp = ''
    flag = 0
    lastfilepath = ""

    filename_txt = file_name(file_path, '.txt')
    if not filename_txt:
        return
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '测试结果'

    # 节点去重：记录(boardID, Config_Name)对应的行号，重复时覆盖
    node_row_map = {}  # {(boardID, Config_Name): row_number}

    cnt = 0
    for txt_file in filename_txt:
        # 从路径反向解析：文件名\工具名\索引\板ID\路径
        try:
            filename, toolname, indx, boardID, filepath = txt_file[::-1].split('\\', 4)
            lastfilepath = filepath
        except ValueError:
            # 路径层级不足时使用上次的路径
            filename, toolname, indx, boardID = txt_file[::-1].split('\\', 3)
            filepath = lastfilepath

        toolname = toolname[::-1]
        boardID = "boardID:" + boardID[::-1]
        toolname_new = toolname

        col = 3

        # 处理CDM、H2和HD_Tune工具文件（HD_Tune包含TLC性能/功耗摘要数据）
        if any(tool in toolname_new for tool in ["CDM", "H2", "HD_Tune"]):
            try:
                # 第一遍扫描：提取Config_Name用于去重判断
                config_name_for_check = toolname
                with open(txt_file, 'rb') as f:
                    for line in f:
                        try:
                            get_buf = line.decode('utf-8')
                        except UnicodeDecodeError:
                            temp = ''
                            for bchar in line:
                                try:
                                    temp += format(bchar, 'c')
                                except UnicodeDecodeError:
                                    temp += ' '
                            get_buf = temp

                        if re.match("^Config_Name:", get_buf):
                            tool_configename = get_buf.split('.', 1)
                            config_name_for_check = tool_configename[0]
                            break

                # 去重检查：如果重复则覆盖之前的行，否则新增行
                node_key = (boardID, config_name_for_check)
                if node_key in node_row_map:
                    # 重复节点，使用之前的行号，覆盖旧数据
                    prev_cnt = cnt  # 保存当前计数器，避免覆盖后影响后续新节点
                    cnt = node_row_map[node_key]
                    print(f"覆盖重复节点（保留最后一份）: {boardID} - {config_name_for_check} [行{cnt}]")
                else:
                    # 新节点，分配新行号
                    prev_cnt = None
                    cnt += 1
                    node_row_map[node_key] = cnt

                # 第二遍扫描：正式提取数据
                with open(txt_file, 'rb') as f:
                    for line in f:
                        # 尝试UTF-8解码，失败则逐字符转换
                        try:
                            get_buf = line.decode('utf-8')
                        except UnicodeDecodeError:
                            temp = ''
                            for bchar in line:
                                try:
                                    temp += format(bchar, 'c')
                                except UnicodeDecodeError:
                                    temp += ' '
                            get_buf = temp

                        # 提取Config_Name作为工具配置名称
                        if re.match("^Config_Name:", get_buf):
                            if flag:
                                toolname_new = toolname
                            tool_configename = get_buf.split('.', 1)
                            toolname_new = tool_configename[0]
                            flag = 1

                        # 匹配关键字并提取数据
                        for key in keys:
                            if re.match('^' + key, get_buf):
                                if ',' in get_buf:
                                    get_buf = get_buf.split(",", 1)[0]  # 取逗号前的值
                                get_buf = get_buf.strip()
                                col = save_board_error_info(worksheet, cnt, col, get_buf)
            except Exception as e:
                print(f"处理文件{txt_file}时出错: {e}")
                continue

            worksheet.cell(cnt, 1).value = boardID
            worksheet.cell(cnt, 2).value = toolname_new

            # 恢复计数器：覆盖操作后恢复原来的计数器值，避免下一个新节点行号错误
            if prev_cnt is not None:
                cnt = prev_cnt

    # 保存为Excel并导出CSV供后续处理
    excel_path = "Performance.xlsx"
    workbook.save(excel_path)
    df = pd.read_excel(excel_path)
    csv_file = 'Performance.csv'
    df.to_csv(csv_file, index=False)


def extract_key_value_pairs(input_string):
    """提取键值对（格式：key1:value1,key2:value2）"""
    pairs = input_string.split(',')
    result = {}
    for pair in pairs:
        key_value = pair.split(':')
        if len(key_value) == 2:
            result[key_value[0]] = key_value[1]
    return result


def extract_key(key_value_pairs, key):
    """从键值对中提取指定键的值"""
    try:
        return str(key_value_pairs[key])
    except KeyError:
        return None


def write_cell(sheet, row, col, value):
    """安全写入单元格，尝试转为float"""
    if value is None or value == '' or value == 'NULL':
        sheet.cell(row, col, None)
    else:
        try:
            sheet.cell(row, col, float(value))
        except (ValueError, TypeError):
            sheet.cell(row, col, str(value))


def Write_To_EVT_Report_0(FILE, filename, Save_In_Work_Directory):
    """将Performance.csv中的数据按配置类型分发到EVT报告的不同工作表和位置"""
    template_file = r"Performance.csv"
    print(3333, filename)
    wb = load_workbook(filename)
    wb1 = wb["性能数据"]
    wb2 = wb["功耗测试"]

    # 定义数据字段映射（顺序与Excel列对应）
    Sleep_Power_Keys = ['Sleep_Cur_VCC_Avg', 'Sleep_Cur_VCCQ_Avg']
    CMD_Performace_Keys = ['Max_Seq_Rd_Spd', 'Max_Rand_Rd_Spd', 'Max_Seq_Wr_Spd', 'Max_Rand_Wr_Spd']
    CMD_Power_VCC_Keys = ['Seq_Rd_Cur_VCC_Avg', 'Rand_Rd_Cur_VCC_Avg', 'Seq_Wr_Cur_VCC_Avg', 'Rand_Wr_Cur_VCC_Avg',
                          'Sleep_Cur_VCC_Avg']
    CMD_Power_VCCQ_Keys = ['Seq_Rd_Cur_VCCQ_Avg', 'Rand_Rd_Cur_VCCQ_Avg', 'Seq_Wr_Cur_VCCQ_Avg', 'Rand_Wr_Cur_VCCQ_Avg',
                           'Sleep_Cur_VCCQ_Avg']
    TLC_Performace_Keys = ['TLC_Seq_Read', 'TLC_Rand_Read', 'TLC_Seq_Write', 'TLC_Seq_Write_GC', 'AveWriteSpeed']
    TLC_VCC_Keys = ['TLC_SeqRead_Cur_VCC', 'TLC_RandRead_Cur_VCC', 'TLC_SeqWrite_Cur_VCC']
    TLC_VCCQ_Keys = ['TLC_SeqRead_Cur_VCCQ', 'TLC_RandRead_Cur_VCCQ', 'TLC_SeqWrite_Cur_VCCQ']

    # 按boardID分配列：同一节点的所有测试数据写入同一列，不同节点写入不同列
    _board_col = {}  # boardID -> column
    _next_col = [10]  # 使用列表以支持闭包修改

    def _get_col(bid):
        """获取写入列号：相同boardID始终返回同一列，不同boardID分配新列"""
        if bid in _board_col:
            return _board_col[bid]
        c = _next_col[0]
        _next_col[0] += 1
        _board_col[bid] = c
        return c

    # 各配置类型在工作表中的起始行
    rows = {
        'HS400': 3, 'HS200': 4, 'DDR50': 5, 'SDR50': 6, 'SDR26': 7,
        'HD_IDA': 13, 'HD_CLEAN': 25,
    }

    with open(template_file, 'r', encoding='utf-8') as f:
        for line in f:
            key_value_pairs = extract_key_value_pairs(line)
            config = extract_key(key_value_pairs, 'Config_Name')

            # HS400配置：写入性能和功耗数据
            if config == "IDA-CMDQ-HS400":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb1, 2, c, bid)
                write_cell(wb2, 2, c, bid)

                r0 = rows['HS400']
                write_cell(wb1, r0, c, extract_key(key_value_pairs, CMD_Performace_Keys[0]))
                write_cell(wb1, r0 + 5, c, extract_key(key_value_pairs, CMD_Performace_Keys[1]))
                write_cell(wb1, r0 + 12, c, extract_key(key_value_pairs, CMD_Performace_Keys[2]))
                write_cell(wb1, r0 + 17, c, extract_key(key_value_pairs, CMD_Performace_Keys[3]))

                write_cell(wb2, r0, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[0]))
                write_cell(wb2, r0 + 10, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[1]))  # 5*2=10
                write_cell(wb2, r0 + 24, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[2]))  # (5+7)*2=24
                write_cell(wb2, r0 + 34, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[3]))  # (5+7+5)*2=34

                write_cell(wb2, r0 + 1, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[0]))
                write_cell(wb2, r0 + 11, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[1]))
                write_cell(wb2, r0 + 25, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[2]))
                write_cell(wb2, r0 + 35, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[3]))

                write_cell(wb2, 49, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 50, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))

            # HS200配置
            elif config == "IDA-CMDQ-HS200":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                r0 = rows['HS200']
                write_cell(wb1, r0, c, extract_key(key_value_pairs, CMD_Performace_Keys[0]))
                write_cell(wb1, r0 + 5, c, extract_key(key_value_pairs, CMD_Performace_Keys[1]))
                write_cell(wb1, r0 + 12, c, extract_key(key_value_pairs, CMD_Performace_Keys[2]))
                write_cell(wb1, r0 + 17, c, extract_key(key_value_pairs, CMD_Performace_Keys[3]))

                write_cell(wb2, r0 + 1, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[0]))
                write_cell(wb2, r0 + 11, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[1]))
                write_cell(wb2, r0 + 25, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[2]))
                write_cell(wb2, r0 + 35, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[3]))

                write_cell(wb2, r0 + 2, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[0]))
                write_cell(wb2, r0 + 12, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[1]))
                write_cell(wb2, r0 + 26, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[2]))
                write_cell(wb2, r0 + 36, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[3]))

                write_cell(wb2, 51, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 52, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))

            # DDR50配置
            elif config == "IDA-CMDQ-DDR50":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                r0 = rows['DDR50']
                write_cell(wb1, r0, c, extract_key(key_value_pairs, CMD_Performace_Keys[0]))
                write_cell(wb1, r0 + 5, c, extract_key(key_value_pairs, CMD_Performace_Keys[1]))
                write_cell(wb1, r0 + 12, c, extract_key(key_value_pairs, CMD_Performace_Keys[2]))
                write_cell(wb1, r0 + 17, c, extract_key(key_value_pairs, CMD_Performace_Keys[3]))

                write_cell(wb2, r0 + 2, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[0]))
                write_cell(wb2, r0 + 12, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[1]))
                write_cell(wb2, r0 + 26, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[2]))
                write_cell(wb2, r0 + 36, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[3]))

                write_cell(wb2, r0 + 3, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[0]))
                write_cell(wb2, r0 + 13, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[1]))
                write_cell(wb2, r0 + 27, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[2]))
                write_cell(wb2, r0 + 37, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[3]))

            # SDR50配置
            elif config == "IDA-CMDQ-SDR50":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                r0 = rows['SDR50']
                write_cell(wb1, r0, c, extract_key(key_value_pairs, CMD_Performace_Keys[0]))
                write_cell(wb1, r0 + 5, c, extract_key(key_value_pairs, CMD_Performace_Keys[1]))
                write_cell(wb1, r0 + 12, c, extract_key(key_value_pairs, CMD_Performace_Keys[2]))
                write_cell(wb1, r0 + 17, c, extract_key(key_value_pairs, CMD_Performace_Keys[3]))

                write_cell(wb2, r0 + 3, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[0]))
                write_cell(wb2, r0 + 13, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[1]))
                write_cell(wb2, r0 + 27, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[2]))
                write_cell(wb2, r0 + 37, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[3]))

                write_cell(wb2, r0 + 4, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[0]))
                write_cell(wb2, r0 + 14, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[1]))
                write_cell(wb2, r0 + 28, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[2]))
                write_cell(wb2, r0 + 38, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[3]))

            # SDR26配置
            elif config == "IDA-CMDQ-SDR26":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                r0 = rows['SDR26']
                write_cell(wb1, r0, c, extract_key(key_value_pairs, CMD_Performace_Keys[0]))
                write_cell(wb1, r0 + 5, c, extract_key(key_value_pairs, CMD_Performace_Keys[1]))
                write_cell(wb1, r0 + 12, c, extract_key(key_value_pairs, CMD_Performace_Keys[2]))
                write_cell(wb1, r0 + 17, c, extract_key(key_value_pairs, CMD_Performace_Keys[3]))

                write_cell(wb2, r0 + 4, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[0]))
                write_cell(wb2, r0 + 14, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[1]))
                write_cell(wb2, r0 + 28, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[2]))
                write_cell(wb2, r0 + 38, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[3]))

                write_cell(wb2, r0 + 5, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[0]))
                write_cell(wb2, r0 + 15, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[1]))
                write_cell(wb2, r0 + 29, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[2]))
                write_cell(wb2, r0 + 39, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[3]))

            # HD Tune IDA配置：TLC性能测试
            elif config in ["CA-HD Tune-IDA", "HD Tune-IDA"]:
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb1, rows['HD_IDA'], c, extract_key(key_value_pairs, TLC_Performace_Keys[0]))
                write_cell(wb1, rows['HD_IDA'] + 1, c, extract_key(key_value_pairs, TLC_Performace_Keys[1]))
                write_cell(wb1, 26, c, extract_key(key_value_pairs, TLC_Performace_Keys[3]))
                write_cell(wb1, 27, c, extract_key(key_value_pairs, TLC_Performace_Keys[4]))

                write_cell(wb2, 23, c, extract_key(key_value_pairs, TLC_VCC_Keys[0]))
                write_cell(wb2, 24, c, extract_key(key_value_pairs, TLC_VCCQ_Keys[0]))
                write_cell(wb2, 25, c, extract_key(key_value_pairs, TLC_VCC_Keys[1]))
                write_cell(wb2, 26, c, extract_key(key_value_pairs, TLC_VCCQ_Keys[1]))
                write_cell(wb2, 59, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 60, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))

            # HD Tune CLEAN配置
            elif config in ["CA-HD Tune-CLEAN", "HD Tune-CLEAN"]:
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb1, rows['HD_CLEAN'], c, extract_key(key_value_pairs, TLC_Performace_Keys[2]))
                write_cell(wb2, 47, c, extract_key(key_value_pairs, TLC_VCC_Keys[0]))
                write_cell(wb2, 48, c, extract_key(key_value_pairs, TLC_VCCQ_Keys[0]))

            # Sleep功耗测试（无数据状态）
            elif config == "SLC-SleepCur-WithoutData-DDR50":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb2, 53, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 54, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))
            elif config == "SLC-SleepCur-WithoutData-SDR50":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb2, 55, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 56, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))
            elif config == "SLC-SleepCur-WithoutData-SDR26":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb2, 57, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 58, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))

            # 高温(HT)环境测试配置
            elif config in ["CA-HD_Tune-IDA-HT", "HD Tune-IDA-HT"]:
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb1, 30, c, extract_key(key_value_pairs, TLC_Performace_Keys[0]))
                write_cell(wb1, 31, c, extract_key(key_value_pairs, TLC_Performace_Keys[1]))
                write_cell(wb2, 65, c, extract_key(key_value_pairs, TLC_VCC_Keys[0]))
                write_cell(wb2, 66, c, extract_key(key_value_pairs, TLC_VCCQ_Keys[0]))
                write_cell(wb2, 67, c, extract_key(key_value_pairs, TLC_VCC_Keys[1]))
                write_cell(wb2, 68, c, extract_key(key_value_pairs, TLC_VCCQ_Keys[1]))
                write_cell(wb2, 85, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 86, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))

            elif config == "IDA-CMDQ-HS400-HT":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb1, 28, c, extract_key(key_value_pairs, CMD_Performace_Keys[0]))
                write_cell(wb1, 29, c, extract_key(key_value_pairs, CMD_Performace_Keys[1]))
                write_cell(wb1, 32, c, extract_key(key_value_pairs, CMD_Performace_Keys[2]))
                write_cell(wb1, 33, c, extract_key(key_value_pairs, CMD_Performace_Keys[3]))

                write_cell(wb2, 61, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[0]))
                write_cell(wb2, 63, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[1]))
                write_cell(wb2, 69, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[2]))
                write_cell(wb2, 71, c, extract_key(key_value_pairs, CMD_Power_VCC_Keys[3]))

                write_cell(wb2, 62, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[0]))
                write_cell(wb2, 64, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[1]))
                write_cell(wb2, 70, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[2]))
                write_cell(wb2, 72, c, extract_key(key_value_pairs, CMD_Power_VCCQ_Keys[3]))

                write_cell(wb2, 75, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 76, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))

            elif config in ["CA-HD_Tune-CLEAN-HT", "HD Tune-CLEAN-HT"]:
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb1, 34, c, extract_key(key_value_pairs, TLC_Performace_Keys[2]))
                write_cell(wb2, 73, c, extract_key(key_value_pairs, TLC_VCC_Keys[0]))
                write_cell(wb2, 74, c, extract_key(key_value_pairs, TLC_VCCQ_Keys[0]))

            elif config == "SLC-SleepCur-WithData-HS200-HT":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb2, 77, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 78, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))
            elif config == "SLC-SleepCur-WithoutData-DDR50-HT":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb2, 79, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 80, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))
            elif config == "SLC-SleepCur-WithoutData-SDR50-HT":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb2, 81, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 82, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))
            elif config == "SLC-SleepCur-WithoutData-SDR26-HT":
                bid = extract_key(key_value_pairs, "boardID")
                c = _get_col(bid)
                write_cell(wb2, 83, c, extract_key(key_value_pairs, Sleep_Power_Keys[0]))
                write_cell(wb2, 84, c, extract_key(key_value_pairs, Sleep_Power_Keys[1]))

    # 保存到目标文件夹
    folder_name = FILE.rpartition('\\')[2]
    output_filename = os.path.basename(filename)
    output_path = os.path.join(folder_name, output_filename)

    print("Saving to:", output_path)
    wb.save(output_path)


def CDM_Speed_Collect(FILE, filename, Save_In_Work_Directory):
    """主流程：提取关键字段、写入EVT报告（觧压已在主模块完成）"""
    # 定义需要从测试日志中提取的数据字段
    keys = [
        # CDM性能测试
        'Max_Seq_Rd_Spd:', 'Max_Seq_Wr_Spd:', 'Max_Rand_Wr_Spd:', 'Max_Rand_Rd_Spd:',
        # 功耗测试
        'Sleep_Cur_VCC_Avg', 'Sleep_Cur_VCCQ_Avg',
        'Seq_Wr_Cur_VCCQ_Avg', 'Seq_Wr_Cur_VCC_Avg',
        'Seq_Rd_Cur_VCCQ_Avg', 'Seq_Rd_Cur_VCC_Avg',
        'Rand_Wr_Cur_VCCQ_Avg', 'Rand_Wr_Cur_VCC_Avg',
        'Rand_Rd_Cur_VCCQ_Avg', 'Rand_Rd_Cur_VCC_Avg',
        # HD TUNE性能测试
        'AveWriteSpeed', 'TLC_Seq_Write_GC', 'TLC_Seq_Write', 'TLC_Seq_Read', 'TLC_Rand_Read',
        # HD TUNE功耗测试
        'TLC_SeqRead_Cur_VCC', 'TLC_SeqRead_Cur_VCCQ',
        'TLC_SeqWrite_Cur_VCC', 'TLC_SeqWrite_Cur_VCCQ',
        'TLC_SeqWriteGC_Cur_VCC', 'TLC_SeqWriteCG_Cur_VCCQ',
        'TLC_RandRead_Cur_VCC'
    ]

    print(f"关键字数量: {len(keys)}")
    test_start = time.time()

    process_fields(FILE, keys)
    Write_To_EVT_Report_0(FILE, filename, Save_In_Work_Directory)

    test_end = time.time()
    print(f'程序运行时间为：{test_end - test_start:.2f} 秒')


EXCEL_TEMPLATE = "竞品分析基本测试项.xlsx"

if __name__ == '__main__':
    Pop_Up_Window_1()
    file_path = os.path.join(FILE, EXCEL_TEMPLATE)
    print(FILE, file_path)
    CDM_Speed_Collect(FILE, file_path, False)
    messagebox.showinfo('提示', '确认完成')
