import os
import tempfile
import shutil
import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
import json
import sys
from openpyxl.utils import get_column_letter

# ✅ 新增：支持多线程和消息队列
import threading
import queue

# 核心提取逻辑已抽离至 extraction_core（纯逻辑，便于单元测试）
from extraction_core import (
    extract_content_after_keyword,
    extract_nested_zips,
    get_config_name,
    build_wide_rows,
)

# 配置与路径设置
def get_script_dir():
    """获取程序所在目录，兼容 PyInstaller 打包"""
    if getattr(sys, 'frozen', False):
        # 打包模式：sys.executable 是 .exe 文件路径
        return os.path.dirname(sys.executable)
    else:
        # 开发模式：__file__ 是 .py 文件路径
        return os.path.dirname(os.path.abspath(__file__))

SCRIPT_DIR = get_script_dir()

# 版本号唯一事实源：窗口标题与完成提示均引用此处，避免多处不同步
APP_VERSION = "1.4"

DEFAULT_TEMPLATE = os.path.join(SCRIPT_DIR, "template.xlsx")
DEFAULT_CONFIG = os.path.join(SCRIPT_DIR, "config.json")

# ========================
# 全局控制变量（新增，不影响原有逻辑）
# ========================
processing = False
message_queue = queue.Queue()


# ========================
# GUI 全局变量：提取模式开关
# ========================
extract_all_mode = None  # 将在 GUI 创建时初始化
wide_table_mode = None   # 将在 GUI 创建时初始化（宽表模式开关）


# ========================
# 核心执行函数（原 run_script 重命名为 run_script_in_thread，逻辑未改，仅增加队列反馈）
# ========================

def run_script_in_thread():
    """主执行函数（支持新格式 tools）—— 在后台线程中运行"""
    global processing

    folder_path = folder_entry.get().strip()
    excel_path = excel_entry.get().strip()
    config_path = config_entry.get().strip()

    # 参数校验
    if not folder_path:
        message_queue.put(("error", "请先选择包含ZIP文件的文件夹！"))
        processing = False
        return
    if not os.path.exists(folder_path):
        message_queue.put(("error", f"文件夹不存在：\n{folder_path}"))
        processing = False
        return
    if not excel_path or not os.path.exists(excel_path):
        message_queue.put(("error", f"Excel模板文件不存在：\n{excel_path}"))
        processing = False
        return
    if not config_path or not os.path.exists(config_path):
        message_queue.put(("error", f"配置文件不存在：\n{config_path}"))
        processing = False
        return

    try:
        # === 读取新格式配置：tools 数组 ===
        with open(config_path, 'r', encoding='utf-8') as f:
            config = json.load(f)
        sheet_name = config.get("sheet_name", "Sheet1")
        tools = config.get("tools", [])

        if not tools:
            message_queue.put(("warning", "配置文件中未定义 tools！"))
            processing = False
            return

        # 收集所有关键字（用于“检测实际存在的关键字”阶段）
        all_keywords = set()
        for tool in tools:
            keywords = tool.get("keyword_to_column", {}).keys()
            all_keywords.update(keywords)
        all_keywords = list(all_keywords)

        if not all_keywords:
            message_queue.put(("warning", "配置文件中未定义任何关键字！"))
            processing = False
            return

        # === 获取用户选择的模式 ===
        use_wide_table = wide_table_mode.get()
        # 宽表模式下每个节点对同一关键字只保留一个结果，强制仅取最后一次
        use_extract_all = extract_all_mode.get() and not use_wide_table

        # === 步骤1：递归解压 ZIP ===
        message_queue.put(("info", "🔄 正在解压 ZIP 文件..."))
        extract_nested_zips(folder_path)

        # === 步骤2：收集所有 TXT 文件 ===
        txt_files = []
        for root, _, files in os.walk(folder_path):
            for file in files:
                if file.lower().endswith('.txt'):
                    txt_files.append(os.path.join(root, file))

        if not txt_files:
            message_queue.put(("warning", "未找到任何 TXT 文件！\n请确认 ZIP 已正确解压。"))
            processing = False
            return

        total_files = len(txt_files)
        message_queue.put(("info", f"📄 共发现 {total_files} 个 TXT 文件，开始处理..."))

        # === 步骤3：检测实际存在的关键字 ===
        present_keywords = []
        for keyword in all_keywords:
            for txt_file in txt_files:
                if extract_content_after_keyword(txt_file, keyword, last_only=True) is not None:
                    present_keywords.append(keyword)
                    break

        if not present_keywords:
            message_queue.put(("warning", "未在任何 TXT 文件中找到配置的关键字！"))
            processing = False
            return

        # === 步骤4：生成 Excel 汇总表 ===
        summary_excel_path = os.path.join(folder_path, "Summary.xlsx")
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        temp_file.close()

        # 复制模板
        shutil.copy(excel_path, temp_file.name)
        wb = openpyxl.load_workbook(temp_file.name)

        if sheet_name not in wb.sheetnames:
            available = "、".join(wb.sheetnames)
            raise ValueError(f"模板中未找到工作表 '{sheet_name}'，可用：{available}")
        ws = wb[sheet_name]

        # === 写表头 ===
        if use_wide_table:
            # 宽表：固定列 Script/配置名/关键字，板子号列在聚合完成后从 D 列起动态写入
            ws.cell(row=1, column=1, value="Script")
            ws.cell(row=1, column=2, value="配置名")
            ws.cell(row=1, column=3, value="关键字")
        else:
            ws.cell(row=1, column=1, value="板子号")
            ws.cell(row=1, column=2, value="Script")
            ws.cell(row=1, column=3, value="Config_Name")
            ws.cell(row=1, column=4, value="关键字")
            ws.cell(row=1, column=5, value="数据")

        # === 写数据：按 Config_Name 匹配 tool，提取其关键字 ===
        row_idx = 2  # 长表从第2行开始
        processed_count = 0
        wide_records = []  # 宽表模式收集 (board_id, script, config, keyword, value) 长记录

        for txt_file in txt_files:
            processed_count += 1
            if processed_count % 5 == 0 or processed_count == total_files:
                message_queue.put(("progress", f"⏳ 正在处理 {processed_count}/{total_files}"))

            try:
                # 获取相对路径，用于判断层级
                rel_path = os.path.relpath(txt_file, folder_path)
                parts = rel_path.replace('\\', '/').split('/')
                is_deep_structure = len(parts) >= 3

                # --- 提取 board_id ---
                if is_deep_structure:
                    board_id = parts[0]
                else:
                    filename = os.path.basename(txt_file)
                    name_no_ext = os.path.splitext(filename)[0]
                    board_id = name_no_ext.split('_', 1)[0]

                # --- 提取 script_value ---
                if is_deep_structure:
                    script_value = os.path.basename(os.path.dirname(txt_file))
                else:
                    script_value = extract_content_after_keyword(txt_file, "Script:") or ""

                # --- 提取 config_name_value ---
                config_name_value = get_config_name(txt_file)

                # --- 查找匹配的 tool（支持全局工具）---
                matched_tools = [
                    tool for tool in tools
                    if tool.get("name") == config_name_value or not tool.get("name")
                ]

                if not matched_tools:
                    print(f"【跳过】{os.path.basename(txt_file)} - 未找到匹配 Config_Name '{config_name_value}' 的 tool")
                    continue

                # --- 提取每个匹配 tool 的所有关键字 ---
                for tool in matched_tools:
                    keyword_to_column = tool.get("keyword_to_column", {})
                    for keyword in keyword_to_column.keys():
                        # 跳过保留关键字
                        if keyword.lower() in ["script", "config_name"]:
                            continue

                        # 如果该关键字不在“实际存在”的列表中，跳过（可选优化）
                        if keyword not in present_keywords:
                            continue

                        # ✅ 根据用户选择决定是否提取全部（宽表模式下强制只取最后一次）
                        results = extract_content_after_keyword(txt_file, keyword, last_only=not use_extract_all)

                        if not results:
                            print(f"【未找到关键字】文件路径: {txt_file}, 关键字: {keyword}")
                            continue

                        if use_wide_table:
                            # 宽表：只收记录，表头/单元格在全部文件处理完后统一写入
                            last_value = results if isinstance(results, str) else results[-1]
                            wide_records.append((board_id, script_value, config_name_value, keyword.strip(), last_value))
                            continue

                        # 统一转为列表处理
                        values_to_write = results if isinstance(results, list) else [results]

                        for value in values_to_write:
                            ws.cell(row=row_idx, column=1, value=board_id)
                            ws.cell(row=row_idx, column=2, value=script_value)
                            ws.cell(row=row_idx, column=3, value=config_name_value)
                            ws.cell(row=row_idx, column=4, value=keyword.strip())
                            ws.cell(row=row_idx, column=5, value=value)
                            row_idx += 1

            except Exception as e:
                print(f"【跳过】处理文件失败 {txt_file}: {e}")
                continue

        # === 宽表模式：聚合长记录，动态写入板子号列头与单元格 ===
        if use_wide_table:
            board_columns, wide_rows = build_wide_rows(wide_records)
            for i, board in enumerate(board_columns):
                ws.cell(row=1, column=4 + i, value=board)
            for row_key, row_data in enumerate(wide_rows, start=2):
                ws.cell(row=row_key, column=1, value=row_data["script"])
                ws.cell(row=row_key, column=2, value=row_data["config"])
                ws.cell(row=row_key, column=3, value=row_data["keyword"])
                for j, board in enumerate(board_columns):
                    value = row_data["values"].get(board)
                    if value is not None:
                        ws.cell(row=row_key, column=4 + j, value=value)
            row_idx = len(wide_rows) + 1  # 用于完成提示的行数统计

        # === 自动调整列宽 ===
        for col in range(1, ws.max_column + 1):
            max_length = 10
            col_letter = get_column_letter(col)
            for cell in ws[col_letter]:
                if cell.value:
                    value_str = str(cell.value)
                    length = sum(2 if ord(c) > 255 else 1 for c in value_str)
                    max_length = max(max_length, length)
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

        # === 统一设置行高 ===
        for row in range(1, ws.max_row + 1):
            ws.row_dimensions[row].height = 18

        # 保存结果
        wb.save(temp_file.name)
        shutil.copy(temp_file.name, summary_excel_path)
        os.remove(temp_file.name)

        # 完成提示（仅弹窗，不自动打开文件夹）
        if use_wide_table:
            msg = (f"成功生成 {row_idx - 1} 行宽表结果（{len(board_columns)} 个板子号列）！\n"
                   f"结果已保存至：\n{summary_excel_path}")
        else:
            msg = f"成功处理 {row_idx - 1} 个数据行！\n结果已保存至：\n{summary_excel_path}"
        print(f"【完成】{msg}")
        message_queue.put(("success", msg))

    except Exception as e:
        error_msg = f"处理过程中发生错误：\n{str(e)}"
        print(f"【异常】{error_msg}")
        message_queue.put(("error", error_msg))
    finally:
        processing = False
        message_queue.put(("enable_buttons", None))


# ========================
# GUI 控制函数（精简版）
# ========================

def run_script():
    """启动后台线程执行任务"""
    global processing
    if processing:
        messagebox.showwarning("警告", "任务已在运行中，请稍候...")
        return

    processing = True
    disable_buttons()
    threading.Thread(target=run_script_in_thread, daemon=True).start()
    root.after(100, check_message_queue)


def check_message_queue():
    """定期检查队列中的消息并更新 GUI"""
    try:
        while True:
            msg_type, content = message_queue.get_nowait()
            if msg_type == "error":
                messagebox.showerror("错误", content)
            elif msg_type == "warning":
                messagebox.showwarning("警告", content)
            elif msg_type == "info" or msg_type == "progress":
                status_label.config(text=content)
            elif msg_type == "success":
                messagebox.showinfo("成功", content)  # 👈 仅弹窗，不自动打开文件夹
            elif msg_type == "enable_buttons":
                enable_buttons()
    except queue.Empty:
        pass
    # 继续监听
    if processing:
        root.after(100, check_message_queue)


def disable_buttons():
    """禁用输入框和开始按钮"""
    folder_entry.config(state="disabled")
    excel_entry.config(state="disabled")
    config_entry.config(state="disabled")
    for widget in root.winfo_children():
        if isinstance(widget, tk.Button) and widget.cget("text") == "开始提取":
            widget.config(state="disabled")


def enable_buttons():
    """恢复输入框和按钮"""
    folder_entry.config(state="normal")
    excel_entry.config(state="normal")
    config_entry.config(state="normal")
    for widget in root.winfo_children():
        if isinstance(widget, tk.Button) and widget.cget("text") == "开始提取":
            widget.config(state="normal")
    status_label.config(text="就绪")


# ========================
# GUI 界面函数（未改动）
# ========================

def select_folder():
    path = filedialog.askdirectory(title="选择包含ZIP文件的文件夹")
    if path:
        folder_entry.delete(0, tk.END)
        folder_entry.insert(0, path)


def select_excel_file():
    path = filedialog.askopenfilename(
        title="选择Excel模板文件",
        filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")]
    )
    if path:
        excel_entry.delete(0, tk.END)
        excel_entry.insert(0, path)


def select_config_file():
    path = filedialog.askopenfilename(
        title="选择配置文件",
        filetypes=[("JSON 文件", "*.json"), ("所有文件", "*.*")]
    )
    if path:
        config_entry.delete(0, tk.END)
        config_entry.insert(0, path)


# ========================
# GUI 主界面（已新增复选框）
# ========================

def main():
    """构建并启动 GUI 主界面（封装为函数，使模块可安全导入）"""
    # GUI 控件供其他函数以全局方式引用，保持原有行为不变
    global root, folder_entry, excel_entry, config_entry, status_label, extract_all_mode, wide_table_mode

    root = tk.Tk()
    root.title(f"关键字数据提取器V{APP_VERSION}")
    root.geometry("650x420")  # 增加高度以容纳宽表模式复选框
    root.resizable(False, False)

    # --- 输入区域 ---
    tk.Label(root, text="ZIP文件夹路径：", font=("Arial", 10)).grid(row=0, column=0, padx=15, pady=12, sticky="e")
    folder_entry = tk.Entry(root, width=45, font=("Arial", 10))
    folder_entry.grid(row=0, column=1, padx=10, pady=10)
    tk.Button(root, text="浏览", command=select_folder, font=("Arial", 9)).grid(row=0, column=2, padx=8)

    tk.Label(root, text="Excel模板：", font=("Arial", 10)).grid(row=1, column=0, padx=15, pady=12, sticky="e")
    excel_entry = tk.Entry(root, width=45, font=("Arial", 10))
    excel_entry.grid(row=1, column=1, padx=10, pady=10)
    tk.Button(root, text="浏览", command=select_excel_file, font=("Arial", 9)).grid(row=1, column=2, padx=8)

    tk.Label(root, text="配置文件：", font=("Arial", 10)).grid(row=2, column=0, padx=15, pady=12, sticky="e")
    config_entry = tk.Entry(root, width=45, font=("Arial", 10))
    config_entry.grid(row=2, column=1, padx=10, pady=10)
    tk.Button(root, text="浏览", command=select_config_file, font=("Arial", 9)).grid(row=2, column=2, padx=8)

    # --- 自动填充默认路径 ---
    if os.path.exists(DEFAULT_TEMPLATE):
        excel_entry.insert(0, DEFAULT_TEMPLATE)
    else:
        print(f"【提示】未找到默认模板：{DEFAULT_TEMPLATE}")

    if os.path.exists(DEFAULT_CONFIG):
        config_entry.insert(0, DEFAULT_CONFIG)
    else:
        print(f"【提示】未找到默认配置：{DEFAULT_CONFIG}")

    # --- 提示信息 ---
    tk.Label(root, text="提示：Excel模板和配置json文件将默认从程序目录中寻找，也可指定路径", fg="gray", font=("Arial", 9)).grid(row=3, column=1, sticky="w", pady=(5, 0))

    # --- 新增：输出模式与提取模式选择 ---
    wide_table_mode = tk.BooleanVar(value=False)
    extract_all_mode = tk.BooleanVar(value=False)

    def on_wide_mode_toggle():
        # 宽表模式下每节点只取一个结果，"提取全部匹配项"置灰禁用
        all_match_check.config(state="disabled" if wide_table_mode.get() else "normal")
        if wide_table_mode.get():
            extract_all_mode.set(False)

    tk.Checkbutton(
        root,
        text="宽表模式（按板子号分列，每节点仅取最后一次结果）",
        variable=wide_table_mode,
        command=on_wide_mode_toggle,
        font=("Arial", 9)
    ).grid(row=4, column=1, sticky="w", padx=(80, 0), pady=(0, 2))

    all_match_check = tk.Checkbutton(
        root,
        text="提取全部匹配项（而非仅最后一次）",
        variable=extract_all_mode,
        font=("Arial", 9)
    )
    all_match_check.grid(row=5, column=1, sticky="w", padx=(80, 0), pady=(0, 10))

    # --- 执行按钮 ---
    tk.Button(
        root,
        text="开始提取",
        command=run_script,
        bg="#4CAF50",
        fg="white",
        font=("Arial", 11, "bold"),
        width=15,
        height=1
    ).grid(row=6, column=1, pady=10)

    # --- 状态栏 ---
    status_label = tk.Label(root, text="就绪", fg="blue", font=("Arial", 9))
    status_label.grid(row=7, column=0, columnspan=3, pady=10)

    # 启动 GUI
    root.mainloop()


if __name__ == "__main__":
    main()